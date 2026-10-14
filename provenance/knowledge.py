import hashlib
import io
import json
import logging
import re
import urllib.error
import urllib.request

from . import db
from .config import (
    CHROMA_PATH,
    LLM_API_KEY,
    LLM_BASE_URL,
    LLM_MODEL_NAME,
    LLM_PROVIDER,
    NEO4J_PASSWORD,
    NEO4J_URI,
    NEO4J_USER,
)

LOGGER = logging.getLogger(__name__)

try:
    import chromadb
except Exception:
    chromadb = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    import pymupdf
except Exception:
    pymupdf = None

try:
    from neo4j import GraphDatabase
except Exception:
    GraphDatabase = None

try:
    from langgraph.graph import END, StateGraph
except Exception:
    END = "__end__"
    StateGraph = None


def extract_text(filename, content):
    name = filename.lower()
    if name.endswith((".txt", ".csv")):
        return content.decode("utf-8", errors="replace").strip()
    if name.endswith(".pdf"):
        if pymupdf is None:
            raise ValueError("PDF parsing requires PyMuPDF.")
        document = pymupdf.open(stream=content, filetype="pdf")
        return "\n\n".join(page.get_text() for page in document).strip()
    if name.endswith(".docx"):
        if DocxDocument is None:
            raise ValueError("DOCX parsing requires python-docx.")
        document = DocxDocument(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    raise ValueError("Unsupported file type. Use PDF, DOCX, TXT, or CSV.")


def chunks(text):
    values = [item.strip() for item in re.split(r"\n\s*\n+", text) if item.strip()]
    return values or [text.strip()]


def embedding(text, dimensions=8):
    tokens = re.findall(r"[a-z0-9]{2,}", text.lower())
    return [
        (int(hashlib.sha256((tokens[index % len(tokens)] if tokens else "empty").encode()).hexdigest()[:8], 16) % 1000) / 1000
        for index in range(dimensions)
    ]


def chroma_collection():
    if chromadb is None:
        return None
    client = chromadb.PersistentClient(path=CHROMA_PATH)
    return client.get_or_create_collection("provenance_memory", metadata={"hnsw:space": "cosine"})


def index_chroma(document_id, project_id, title, text):
    collection = chroma_collection()
    if collection is None:
        return
    values = chunks(text)[:20]
    collection.upsert(
        ids=[f"doc-{document_id}-chunk-{index}" for index in range(len(values))],
        documents=values,
        embeddings=[embedding(value) for value in values],
        metadatas=[{"document_id": document_id, "project_id": project_id, "title": title} for _ in values],
    )


def remove_chroma_document(document_id):
    collection = chroma_collection()
    if collection is not None:
        collection.delete(where={"document_id": document_id})


def search_chroma(project_ids, question, limit=3):
    collection = chroma_collection()
    if collection is None or not project_ids:
        return []
    result = collection.query(
        query_embeddings=[embedding(question)],
        n_results=limit,
        where={"project_id": {"$in": project_ids}},
    )
    return result.get("documents", [[]])[0]


def neo4j_driver():
    if GraphDatabase is None or not all((NEO4J_URI, NEO4J_USER, NEO4J_PASSWORD)):
        return None
    return GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))


def sync_graph(project_id, title, text, document_id=None):
    driver = neo4j_driver()
    if driver is None:
        return False
    connection = db.connect()
    project = connection.execute("SELECT name FROM projects WHERE id=?", (project_id,)).fetchone()
    connection.close()
    if not project:
        driver.close()
        return False
    try:
        with driver.session(database="neo4j") as session:
            session.run(
                """MERGE (p:Project {id:$project_id}) SET p.name=$project_name
                MERGE (d:Document {id:$document_id}) SET d.title=$title,d.project_id=$project_id
                MERGE (p)-[:HAS_DOCUMENT]->(d)""",
                project_id=project_id, project_name=project["name"],
                document_id=f"doc-{document_id or f'{project_id}-{title}'}", title=title,
            )
            keywords = list(dict.fromkeys(re.findall(r"[A-Za-z]{4,}", text)))[:10]
            session.run(
                """MATCH (p:Project {id:$project_id}) UNWIND $keywords AS keyword
                MERGE (t:Topic {name:keyword}) MERGE (p)-[:COVERS]->(t)""",
                project_id=project_id, keywords=keywords,
            )
            for decision in decision_candidates(text):
                session.run(
                    """MATCH (d:Document {id:$document_id})
                    MERGE (x:Decision {id:$decision_id})
                    SET x.title=$title,x.status='Proposed',x.evidence=$evidence
                    MERGE (x)-[:SUPPORTED_BY]->(d)""",
                    document_id=f"doc-{document_id or f'{project_id}-{title}'}",
                    decision_id=f"decision-{project_id}-{hashlib.sha256(decision.encode()).hexdigest()[:16]}",
                    title=decision[:120],
                    evidence=title,
                )
    finally:
        driver.close()
    return True


def remove_graph_document(document_id):
    driver = neo4j_driver()
    if driver is None:
        return
    try:
        with driver.session(database="neo4j") as session:
            session.run(
                """MATCH (d:Document {id:$document_id})
                OPTIONAL MATCH (d)<-[:SUPPORTED_BY]-(decision:Decision)
                DETACH DELETE d, decision""",
                document_id=f"doc-{document_id}",
            )
    finally:
        driver.close()


def sync_decisions_graph():
    driver = neo4j_driver()
    if driver is None:
        return 0
    connection = db.connect()
    rows = connection.execute(
        """SELECT d.id,d.project_id,d.title,d.description,d.status,d.evidence,p.name AS project_name
        FROM decisions d JOIN projects p ON p.id=d.project_id"""
    ).fetchall()
    connection.close()
    try:
        with driver.session(database="neo4j") as session:
            for row in rows:
                session.run(
                    """MERGE (p:Project {id:$project_id}) SET p.name=$project_name
                    MERGE (d:Decision {id:$decision_id})
                    SET d.title=$title,d.status=$status,d.evidence=$evidence
                    MERGE (p)-[:HAS_DECISION]->(d)""",
                    project_id=row["project_id"],
                    project_name=row["project_name"],
                    decision_id=f"decision-{row['id']}",
                    title=row["title"],
                    status=row["status"],
                    evidence=row["evidence"] or row["description"],
                )
        return len(rows)
    finally:
        driver.close()


def reindex_existing_documents():
    connection = db.connect()
    rows = connection.execute(
        "SELECT id,project_id,title,content FROM documents WHERE content IS NOT NULL"
    ).fetchall()
    connection.close()
    for row in rows:
        index_chroma(row["id"], row["project_id"], row["title"], row["content"])
        sync_graph(row["project_id"], row["title"], row["content"], row["id"])
    return len(rows), sync_decisions_graph()


def decision_candidates(text):
    terms = r"decision|approved|standardize|selected|preferred|adopt|defer|pause|replace|retire|recommend"
    return list(dict.fromkeys(
        line.strip()[:180] for line in text.splitlines()
        if line.strip() and re.search(rf"(?i)\b({terms})\b", line)
    ))[:5]


def persist_decision_candidates(user_id, project_id, document_id, document_title, text):
    candidates = decision_candidates(text)
    if not candidates:
        return []
    connection = db.connect()
    saved = []
    for candidate in candidates:
        duplicate = connection.execute(
            "SELECT id FROM decisions WHERE project_id=? AND title=? AND evidence LIKE ?",
            (project_id, candidate[:120], f"%document:{document_id}%"),
        ).fetchone()
        if duplicate:
            continue
        cursor = connection.execute(
            """INSERT INTO decisions(project_id,title,description,rationale,alternatives,status,
            proposer_id,owner_id,created_at,updated_at,evidence)
            VALUES(?,?,?,?,?,'Proposed',?,?,?,?,?)""",
            (
                project_id,
                candidate[:120],
                candidate,
                "Extracted from uploaded source; requires human verification.",
                "Not recorded",
                user_id,
                user_id,
                db.now(),
                db.now(),
                f"document:{document_id};title:{document_title};unverified",
            ),
        )
        saved.append(cursor.lastrowid)
    connection.commit()
    connection.close()
    return saved


def process_upload(user, project_id, uploaded_file):
    if not db.can_access_project(user["id"], project_id):
        raise ValueError("You are not authorized to upload documents for this project.")
    extension = uploaded_file.name.rsplit(".", 1)[-1].lower()
    title = uploaded_file.name.rsplit(".", 1)[0]
    connection = db.connect()
    duplicate = connection.execute(
        "SELECT id FROM documents WHERE project_id=? AND title=? AND file_type=?",
        (project_id, title, extension),
    ).fetchone()
    if duplicate:
        connection.close()
        raise ValueError("This document has already been uploaded for this project.")
    text = extract_text(uploaded_file.name, uploaded_file.getvalue())
    if not text:
        connection.close()
        raise ValueError("The uploaded file did not contain readable text.")
    cursor = connection.execute(
        """INSERT INTO documents(project_id,title,file_type,source_type,uploaded_by,uploaded_at,status,version,access_level,content)
        VALUES(?,?,?,'uploaded',?,?,'processed','v1','project',?)""",
        (project_id, title, extension, user["id"], db.now(), text),
    )
    document_id = cursor.lastrowid
    connection.executemany(
        "INSERT INTO document_chunks(document_id,chunk_index,content,created_at) VALUES(?,?,?,?)",
        [(document_id, index, value[:1500], db.now()) for index, value in enumerate(chunks(text))],
    )
    connection.commit()
    connection.close()
    index_chroma(document_id, project_id, title, text)
    sync_graph(project_id, title, text, document_id)
    saved_decisions = persist_decision_candidates(user["id"], project_id, document_id, title, text)
    if saved_decisions:
        db.log_event(user["id"], "decision_candidates_extracted", str(saved_decisions))
    db.log_event(user["id"], "document_upload", f"Uploaded {title}.")
    return document_id


def delete_document(user, document_id):
    record = db.document(document_id)
    if not record:
        raise ValueError("Document not found.")
    if not db.can_access_project(user["id"], record["project_id"]):
        raise ValueError("You are not authorized to delete this document.")
    remove_chroma_document(document_id)
    remove_graph_document(document_id)
    db.delete_document_data(document_id)
    db.log_event(user["id"], "document_delete", f"Deleted document {document_id}.")


def reprocess_document(user, document_id):
    record = db.document(document_id)
    if not record:
        raise ValueError("Document not found.")
    if not db.can_access_project(user["id"], record["project_id"]):
        raise ValueError("You are not authorized to reprocess this document.")
    remove_chroma_document(document_id)
    remove_graph_document(document_id)
    index_chroma(document_id, record["project_id"], record["title"], record["content"] or "")
    sync_graph(record["project_id"], record["title"], record["content"] or "", document_id)
    sync_decisions_graph()
    db.log_event(user["id"], "document_reprocess", f"Reprocessed document {document_id}.")


def search_authorized(user, question, limit=5):
    project_ids = db.authorized_project_ids(user)
    if not question or not project_ids:
        return []
    tokens = set(re.findall(r"[a-z0-9]{3,}", question.lower()))
    connection = db.connect()
    rows = connection.execute(
        f"""SELECT d.*,p.name AS project_name FROM documents d JOIN projects p ON p.id=d.project_id
        WHERE d.project_id IN ({','.join('?' for _ in project_ids)})""",
        project_ids,
    ).fetchall()
    connection.close()
    matches = []
    for row in rows:
        haystack = f"{row['title']} {row['content']}".lower()
        score = sum(haystack.count(token) for token in tokens)
        if score:
            matches.append({"title": row["title"], "project": row["project_name"], "source": row["source_type"], "score": score})
    matches.sort(key=lambda item: item["score"], reverse=True)
    known_text = {row["content"].strip().lower() for row in rows}
    for text in search_chroma(project_ids, question, limit):
        normalized = text.strip().lower()
        if not normalized or normalized in known_text:
            continue
        matches.append({"title": "Semantic memory", "project": "Authorized projects", "source": text[:180], "score": 1})
    unique = []
    seen = set()
    for item in matches:
        key = (item["title"], item["project"], item["source"])
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique[:limit]


def graph_search(user, question, limit=10):
    driver = neo4j_driver()
    project_ids = db.authorized_project_ids(user)
    if driver is None or not project_ids:
        return []
    try:
        with driver.session(database="neo4j") as session:
            result = session.run(
                """MATCH (p:Project)-[*1..2]-(related)
                WHERE p.id IN $project_ids
                  AND (toLower(coalesce(properties(related)['title'],'')) CONTAINS toLower($question)
                    OR toLower(coalesce(properties(related)['name'],'')) CONTAINS toLower($question))
                RETURN labels(related) AS labels, properties(related)['title'] AS title,
                       properties(related)['name'] AS name, p.name AS project
                LIMIT $limit""",
                project_ids=project_ids,
                question=question,
                limit=limit,
            )
            return [record.data() for record in result]
    finally:
        driver.close()


def graph_snapshot(user, project_id=None):
    authorized_ids = db.authorized_project_ids(user)
    if project_id is not None and project_id not in authorized_ids:
        raise ValueError("You are not authorized to inspect this project.")
    project_ids = [project_id] if project_id is not None else authorized_ids
    if not project_ids:
        return {"nodes": [], "edges": [], "backend": "none"}

    driver = neo4j_driver()
    if driver is not None:
        try:
            with driver.session(database="neo4j") as session:
                result = session.run(
                    """MATCH (p:Project)-[r]-(related)
                    WHERE p.id IN $project_ids
                    RETURN p.id AS project_id, p.name AS project_name,
                           labels(related) AS labels, properties(related) AS properties,
                           type(r) AS relationship""",
                    project_ids=project_ids,
                )
                nodes = {}
                edges = []
                for record in result:
                    project_key = f"project:{record['project_id']}"
                    nodes[project_key] = {
                        "id": project_key,
                        "label": record["project_name"],
                        "type": "Project",
                    }
                    properties = record["properties"] or {}
                    related_id = properties.get("id") or properties.get("name") or properties.get("title")
                    if not related_id:
                        continue
                    related_key = f"entity:{record['labels'][0]}:{related_id}"
                    nodes[related_key] = {
                        "id": related_key,
                        "label": properties.get("title") or properties.get("name") or str(related_id),
                        "type": record["labels"][0],
                    }
                    edges.append({"source": project_key, "target": related_key, "label": record["relationship"]})
                return {"nodes": list(nodes.values()), "edges": edges, "backend": "neo4j"}
        finally:
            driver.close()

    connection = db.connect()
    placeholders = ",".join("?" for _ in project_ids)
    projects = connection.execute(
        f"SELECT id,name FROM projects WHERE id IN ({placeholders})", project_ids
    ).fetchall()
    documents = connection.execute(
        f"SELECT id,project_id,title FROM documents WHERE project_id IN ({placeholders})", project_ids
    ).fetchall()
    decisions = connection.execute(
        f"SELECT id,project_id,title FROM decisions WHERE project_id IN ({placeholders})", project_ids
    ).fetchall()
    connection.close()
    nodes = []
    edges = []
    for project in projects:
        project_key = f"project:{project['id']}"
        nodes.append({"id": project_key, "label": project["name"], "type": "Project"})
        for document in [row for row in documents if row["project_id"] == project["id"]]:
            document_key = f"document:{document['id']}"
            nodes.append({"id": document_key, "label": document["title"], "type": "Document"})
            edges.append({"source": project_key, "target": document_key, "label": "HAS_DOCUMENT"})
        for decision in [row for row in decisions if row["project_id"] == project["id"]]:
            decision_key = f"decision:{decision['id']}"
            nodes.append({"id": decision_key, "label": decision["title"], "type": "Decision"})
            edges.append({"source": project_key, "target": decision_key, "label": "HAS_DECISION"})
    return {"nodes": nodes, "edges": edges, "backend": "sqlite"}


def impact_analysis(user, project_id):
    authorized_ids = db.authorized_project_ids(user)
    if project_id not in authorized_ids:
        raise ValueError("You are not authorized to analyze this project.")

    connection = db.connect()
    project = connection.execute("SELECT id,name FROM projects WHERE id=?", (project_id,)).fetchone()
    documents = connection.execute(
        "SELECT id,title,content FROM documents WHERE project_id=?", (project_id,)
    ).fetchall()
    decisions = connection.execute(
        "SELECT id,title,status FROM decisions WHERE project_id=?", (project_id,)
    ).fetchall()
    connection.close()
    if not project:
        raise ValueError("Project not found.")

    snapshot = graph_snapshot(user, project_id)
    connected = len(snapshot["edges"])
    impacted = []
    for node in snapshot["nodes"]:
        if node["type"] == "Project":
            continue
        impacted.append({
            "entity": node["label"],
            "type": node["type"],
            "relationship_count": sum(
                1 for edge in snapshot["edges"] if node["id"] in (edge["source"], edge["target"])
            ),
            "impact": "High" if node["type"] == "Decision" else "Medium",
        })
    return {
        "project": project["name"],
        "backend": snapshot["backend"],
        "documents": len(documents),
        "decisions": len(decisions),
        "relationships": connected,
        "affected_entities": sorted(impacted, key=lambda item: item["relationship_count"], reverse=True),
    }


def call_llm(question, evidence, graph_evidence):
    if not LLM_API_KEY:
        return None
    prompt = (
        "You are the Provenance organizational memory assistant. Answer only from the supplied "
        "evidence. Distinguish documented facts from inference. If evidence is insufficient, say so. "
        "Include concise source citations using source titles.\n\n"
        f"Question: {question}\nEvidence: {json.dumps(evidence)}\n"
        f"Graph evidence: {json.dumps(graph_evidence)}"
    )
    if LLM_PROVIDER == "gemini":
        endpoint = (
            f"{LLM_BASE_URL.rstrip('/')}/models/{LLM_MODEL_NAME}:generateContent"
            f"?key={LLM_API_KEY}"
        )
        payload = json.dumps({
            "contents": [{"role": "user", "parts": [{"text": prompt}]}],
            "generationConfig": {"temperature": 0.1},
        }).encode("utf-8")
        headers = {"Content-Type": "application/json"}
    else:
        endpoint = f"{LLM_BASE_URL.rstrip('/')}/chat/completions"
        payload = json.dumps({
            "model": LLM_MODEL_NAME,
            "messages": [
                {"role": "system", "content": "You are the Provenance organizational memory assistant."},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.1,
        }).encode("utf-8")
        headers = {
            "Authorization": f"Bearer {LLM_API_KEY}",
            "Content-Type": "application/json",
        }
    request = urllib.request.Request(endpoint, data=payload, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            body = json.loads(response.read().decode("utf-8"))
            if LLM_PROVIDER == "gemini":
                return body["candidates"][0]["content"]["parts"][0]["text"].strip()
            return body["choices"][0]["message"]["content"].strip()
    except (urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
        LOGGER.warning("LLM request failed for provider %s: %s", LLM_PROVIDER, exc)
        return None


def run_reasoning_workflow(question, user):
    def retrieve(state):
        return {
            **state,
            "evidence": search_authorized(user, question),
            "graph_evidence": graph_search(user, question),
        }

    def generate(state):
        response = call_llm(question, state["evidence"], state["graph_evidence"])
        if response:
            return {**state, "answer": response}
        if not state["evidence"] and not state["graph_evidence"]:
            return {**state, "answer": "I could not find enough evidence in your authorized sources to answer confidently."}
        evidence = state["evidence"]
        graph_evidence = state["graph_evidence"]
        details = []
        for item in evidence[:3]:
            details.append(
                f"{item['title']} ({item['project']}) is the closest matching authorized source "
                f"(relevance score {item['score']})."
            )
        for item in graph_evidence[:3]:
            label = ", ".join(item.get("labels") or ["graph entity"])
            name = item.get("title") or item.get("name") or "unnamed entity"
            details.append(f"The knowledge graph connects {label} '{name}' to {item.get('project', 'an authorized project')}.")
        return {
            **state,
            "answer": "I found authorized context, but the language model was unavailable. "
            + " ".join(details)
            + " Review the cited sources before treating this as a confirmed answer.",
        }

    if StateGraph is None:
        return generate(retrieve({}))
    workflow = StateGraph(dict)
    workflow.add_node("retrieve", retrieve)
    workflow.add_node("generate", generate)
    workflow.set_entry_point("retrieve")
    workflow.add_edge("retrieve", "generate")
    workflow.add_edge("generate", END)
    result = workflow.compile().invoke({})
    return result


def answer(question, user):
    result = run_reasoning_workflow(question, user)
    matches = result.get("evidence", [])
    graph_evidence = result.get("graph_evidence", [])
    citations = []
    seen = set()
    for item in matches:
        citation = {"title": item["title"], "project": item["project"], "source": item["source"]}
        key = (citation["title"], citation["project"], citation["source"])
        if key not in seen:
            seen.add(key)
            citations.append(citation)
    for item in graph_evidence:
        citation = {
            "title": item.get("title") or item.get("name") or "Graph entity",
            "project": item.get("project", "Authorized projects"),
            "source": "Neo4j knowledge graph",
        }
        key = (citation["title"], citation["project"], citation["source"])
        if key not in seen:
            seen.add(key)
            citations.append(citation)
    return result["answer"], citations
